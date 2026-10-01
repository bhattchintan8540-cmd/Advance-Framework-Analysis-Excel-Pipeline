"""Build the Excel workbook: raw extract, cleaned table, live formulas, charts, QA."""

from __future__ import annotations

from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from pipeline.catalog import (
    CATEGORY_ALIASES,
    CHANNEL_ALIASES,
    REGION_ALIASES,
    SEGMENT_ALIASES,
    STATUS_ALIASES,
)
from pipeline.metrics import CATEGORIES, CHANNELS, REGIONS, STATUSES

NAVY = "1E3A5F"
TEAL = "1F6F64"
GOLD = "B0893E"
CORAL = "A33B32"
SLATE = "5E6A75"
INK = "1C2430"
CREAM = "F7F4EE"
WHITE = "FFFFFF"
PALE = "E7F0EE"
AMBER = "F8F1DE"
RED_SOFT = "F8E6E4"
LINE = "D5D0C6"
SERIES = ["1E3A5F", "1F6F64", "B0893E", "A33B32", "6E7F8D"]

THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
TITLE_FONT = Font(name="Calibri", size=20, bold=True, color=WHITE)
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color=SLATE)
BODY_FONT = Font(name="Calibri", size=11, color=INK)
KPI_FONT = Font(name="Calibri", size=18, bold=True, color=NAVY)
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _paint(cell, fill, font=None, alignment=None, fmt=None, border=None):
    cell.fill = fill
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border


def _criteria(honor_region=True, honor_channel=True, honor_category=True, honor_cancelled=True, extras=()):
    parts = ["CleanData[OrderDate]", '">="&pStart', "CleanData[OrderDate]", '"<="&pEnd']
    if honor_region:
        parts += ["CleanData[Region]", 'IF(pRegion="All","<>",pRegion)']
    if honor_channel:
        parts += ["CleanData[Channel]", 'IF(pChannel="All","<>",pChannel)']
    if honor_category:
        parts += ["CleanData[Category]", 'IF(pCategory="All","<>",pCategory)']
    if honor_cancelled:
        parts += ["CleanData[Status]", 'IF(pIncludeCancelled="Yes","<>","<>Cancelled")']
    for column, expression in extras:
        parts += [f"CleanData[{column}]", expression]
    return parts


def sumifs(value, **kwargs) -> str:
    parts = [f"CleanData[{value}]", *_criteria(**kwargs)]
    return "=SUMIFS(" + ",".join(parts) + ")"


def countifs(**kwargs) -> str:
    return "=COUNTIFS(" + ",".join(_criteria(**kwargs)) + ")"


def averageifs(value, **kwargs) -> str:
    parts = [f"CleanData[{value}]", *_criteria(**kwargs)]
    return "=IFERROR(AVERAGEIFS(" + ",".join(parts) + "),NA())"


def countifs_status(status_expr, **kwargs) -> str:
    kwargs = dict(kwargs)
    kwargs["honor_cancelled"] = False
    parts = _criteria(**kwargs) + ["CleanData[Status]", status_expr]
    return "=COUNTIFS(" + ",".join(parts) + ")"


def _legend(chart):
    chart.legend = Legend()
    chart.legend.position = "b"


def _size(chart, width=12, height=7):
    chart.width = width
    chart.height = height
    chart.style = 10


def _color_series(chart, colors=SERIES):
    for index, series in enumerate(chart.series):
        series.graphicalProperties.solidFill = colors[index % len(colors)]


def _bar(title, anchor_sheet, data_ref, cat_ref, anchor, *, stacked=False, colors=SERIES):
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked" if stacked else "clustered"
    if stacked:
        chart.overlap = 100
    chart.title = title
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cat_ref)
    chart.y_axis.numFmt = '#,##0'
    _legend(chart)
    _size(chart)
    _color_series(chart, colors)
    anchor_sheet.add_chart(chart, anchor)
    return chart


def build_workbook(raw, clean, rejected, cleaning_log, summary, params, path):
    wb = Workbook()
    filters = params["dashboard_filters"]
    _lists(wb)
    _parameters(wb, params, summary)
    _raw(wb, raw)
    _lookup(wb)
    _cleaned(wb, clean)
    _chart_data(wb, summary)
    _analysis(wb, summary)
    _dashboard(wb, filters, params)
    _insights(wb, summary, params, cleaning_log)
    _qa(wb, summary, cleaning_log, clean, raw, rejected)
    _dictionary(wb)
    _cleaning_log(wb, cleaning_log)
    _cover(wb, params, cleaning_log, summary)
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    # Cover was created last but should open first.
    wb.move_sheet("Cover", offset=-len(wb.sheetnames) + 1)
    wb.active = wb["Cover"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _lists(wb):
    ws = wb.active
    ws.title = "Lists"
    columns = {
        "A": ["All", *REGIONS],
        "B": ["All", *CHANNELS],
        "C": ["All", *CATEGORIES],
        "D": ["Yes", "No"],
    }
    for column, values in columns.items():
        for index, value in enumerate(values, start=1):
            ws.cell(index, ord(column) - 64, value)
    wb.defined_names.add(DefinedName(name="lstRegion", attr_text="Lists!$A$1:$A$5"))
    wb.defined_names.add(DefinedName(name="lstChannel", attr_text="Lists!$B$1:$B$4"))
    wb.defined_names.add(DefinedName(name="lstCategory", attr_text="Lists!$C$1:$C$5"))
    wb.defined_names.add(DefinedName(name="lstYesNo", attr_text="Lists!$D$1:$D$2"))
    ws.sheet_state = "hidden"


def _banner(ws, last_col, title, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    _paint(cell, _fill(NAVY), TITLE_FONT, Alignment(vertical="center", horizontal="left", indent=1))
    ws.row_dimensions[1].height = 32
    for col in range(1, last_col + 1):
        ws.cell(1, col).fill = _fill(NAVY)
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        sub = ws.cell(2, 1, subtitle)
        _paint(sub, _fill(TEAL), Font(name="Calibri", size=11, color=WHITE), Alignment(vertical="center", indent=1))
        ws.row_dimensions[2].height = 20
        for col in range(1, last_col + 1):
            ws.cell(2, col).fill = _fill(TEAL)


def _parameters(wb, params, summary):
    ws = wb.create_sheet("Parameters")
    _banner(
        ws,
        6,
        "Model parameters",
        "Dashboard dropdowns are the filters. Thresholds below drive late-shipment and margin flags.",
    )
    rows = [
        (4, "Parameter", "Value", "What it controls", "Limitation"),
        (5, "Late shipment threshold (days)", params["late_shipment_days"], "LateFlag and the on-time KPI. Edit this cell and Excel recalculates.", "Does not rewrite history; it only changes the cutoff."),
        (6, "Low-margin threshold", params["low_margin_threshold"], "Low-margin KPI. Stored as a fraction (0.15 = 15%).", "Compared with each order's margin, not the weighted regional margin."),
        (7, "High-value percentile", params["high_value_percentile"], "Documented percentile used when the pipeline computed the cutoff.", "Changing this cell does not recompute the cutoff. Rerun the pipeline."),
        (8, "High-value cutoff", round(summary["cutoff"], 2), "HighValueFlag. Orders at or above this net revenue count as high value.", "Cutoff is fixed for the saved parameter set and the All/All/All filter. Dropdowns do not rebuild the percentile."),
        (9, "Random seed", params["random_seed"], "Makes the synthetic extract reproducible.", "The extract is synthetic. It is not a client ledger."),
        (10, "Include cancelled (echo)", '=pIncludeCancelled', "Mirrors the dashboard switch.", "Yes includes cancelled orders in sales KPIs. Return rate and on-time still exclude them."),
        (11, "Analysis start (echo)", "=pStart", "Mirrors the dashboard date.", "Rows outside the window stay on Cleaned_Data and drop out of every SUMIFS."),
        (12, "Analysis end (echo)", "=pEnd", "Mirrors the dashboard date.", "Same window rule as the start date."),
    ]
    for row_number, label, value, controls, limit in rows:
        ws.cell(row_number, 2, label)
        ws.cell(row_number, 3, value)
        ws.cell(row_number, 4, controls)
        ws.cell(row_number, 5, limit)
    for col in range(2, 6):
        _paint(ws.cell(4, col), _fill(NAVY), HEADER_FONT, CENTER, border=THIN)
    for row in range(5, 13):
        for col in range(2, 6):
            _paint(ws.cell(row, col), _fill(WHITE), BODY_FONT, WRAP, border=THIN)
        ws.row_dimensions[row].height = 36
    ws["C6"].number_format = "0.00%"
    ws["C7"].number_format = "0%"
    ws["C8"].number_format = '"$"#,##0.00'
    ws["C11"].number_format = "YYYY-MM-DD"
    ws["C12"].number_format = "YYYY-MM-DD"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 62
    ws.column_dimensions["E"].width = 62
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = "B4:E12"
    wb.defined_names.add(DefinedName(name="pLateDays", attr_text="Parameters!$C$5"))
    wb.defined_names.add(DefinedName(name="pLowMargin", attr_text="Parameters!$C$6"))
    wb.defined_names.add(DefinedName(name="pHighCutoff", attr_text="Parameters!$C$8"))
    ws.conditional_formatting.add("C6", CellIsRule(operator="greaterThan", formula=["1"], fill=_fill(RED_SOFT)))
    ws.conditional_formatting.add("C5", CellIsRule(operator="lessThan", formula=["0"], fill=_fill(RED_SOFT)))


def _write_table(ws, frame: pd.DataFrame, table_name: str):
    headers = list(frame.columns)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        _paint(cell, _fill(NAVY), HEADER_FONT, CENTER, border=THIN)
    for row_index, record in enumerate(frame.itertuples(index=False), start=2):
        for col, value in enumerate(record, start=1):
            if value is None or (isinstance(value, float) and pd.isna(value)):
                written = None
            elif isinstance(value, pd.Timestamp):
                written = value.to_pydatetime()
            else:
                written = value
            cell = ws.cell(row_index, col, written)
            cell.font = BODY_FONT
            cell.border = THIN
            header = headers[col - 1]
            if header in {"OrderDate"} and isinstance(written, datetime):
                cell.number_format = "YYYY-MM-DD"
            elif header in {"UnitPrice", "UnitCost", "GrossRevenue", "NetRevenue", "TotalCost"}:
                cell.number_format = '"$"#,##0.00'
            elif header in {"Discount", "Margin"}:
                cell.number_format = "0.00%"
    last_row = max(2, len(frame) + 1)
    last_col = len(headers)
    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(last_col)}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.auto_filter.ref = table.ref
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"
    for col in range(1, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.row_dimensions[1].height = 22
    return last_row


def _raw(wb, raw):
    ws = wb.create_sheet("Raw_Data")
    _write_table(ws, raw, "RawData")
    ws.sheet_properties.tabColor = "A33B32"


def _lookup(wb):
    ws = wb.create_sheet("Lookup_Maps")
    _banner(ws, 4, "Alias maps", "The cleaning stage applies these maps. They are reproduced here so the workbook can be audited without opening the Python code.")
    catalogs = [
        ("Region", REGION_ALIASES),
        ("Channel", CHANNEL_ALIASES),
        ("Category", CATEGORY_ALIASES),
        ("Status", STATUS_ALIASES),
        ("Segment", SEGMENT_ALIASES),
    ]
    row = 4
    for title, mapping in catalogs:
        ws.cell(row, 2, title)
        _paint(ws.cell(row, 2), _fill(TEAL), HEADER_FONT, CENTER)
        _paint(ws.cell(row, 3), _fill(TEAL), HEADER_FONT, CENTER)
        ws.cell(row, 3, "Standard value")
        ws.cell(row, 2, "Alias")
        for offset, (alias, standard) in enumerate(mapping.items(), start=1):
            ws.cell(row + offset, 2, alias)
            ws.cell(row + offset, 3, standard)
            ws.cell(row + offset, 2).font = BODY_FONT
            ws.cell(row + offset, 3).font = BODY_FONT
        row += len(mapping) + 2
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 24
    ws.sheet_view.showGridLines = False


def _cleaned(wb, clean: pd.DataFrame):
    ws = wb.create_sheet("Cleaned_Data")
    base_columns = [
        "OrderID", "OrderDate", "Region", "Channel", "Category", "Product", "Segment",
        "Quantity", "UnitPrice", "Discount", "UnitCost", "ShipDays", "Status", "SalesRep", "QtyAdjusted",
    ]
    formula_columns = [
        "GrossRevenue", "NetRevenue", "TotalCost", "Margin", "Year", "Quarter", "YearMonth",
        "LateFlag", "HighValueFlag", "ReturnFlag",
    ]
    headers = base_columns + formula_columns
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        _paint(cell, _fill(NAVY), HEADER_FONT, CENTER, border=THIN)
    export = clean.copy()
    for row_index, record in enumerate(export.itertuples(index=False), start=2):
        values = {
            "OrderID": record.OrderID,
            "OrderDate": record.OrderDate.to_pydatetime(),
            "Region": record.Region,
            "Channel": record.Channel,
            "Category": record.Category,
            "Product": record.Product,
            "Segment": record.Segment,
            "Quantity": int(record.Quantity),
            "UnitPrice": float(record.UnitPrice),
            "Discount": float(record.Discount),
            "UnitCost": float(record.UnitCost),
            "ShipDays": None if pd.isna(record.ShipDays) else int(record.ShipDays),
            "Status": record.Status,
            "SalesRep": record.SalesRep,
            "QtyAdjusted": record.QtyAdjusted,
        }
        for col, header in enumerate(base_columns, start=1):
            cell = ws.cell(row_index, col, values[header])
            cell.font = BODY_FONT
            cell.border = THIN
        ws.cell(row_index, 2).number_format = "YYYY-MM-DD"
        ws.cell(row_index, 9).number_format = '"$"#,##0.00'
        ws.cell(row_index, 10).number_format = "0.00%"
        ws.cell(row_index, 11).number_format = '"$"#,##0.00'
        formulas = {
            16: f"=H{row_index}*I{row_index}",
            17: f"=P{row_index}*(1-J{row_index})",
            18: f"=H{row_index}*K{row_index}",
            19: f'=IF(Q{row_index}=0,"",(Q{row_index}-R{row_index})/Q{row_index})',
            20: f"=YEAR(B{row_index})",
            21: f'=YEAR(B{row_index})&"-Q"&ROUNDUP(MONTH(B{row_index})/3,0)',
            22: f'=TEXT(B{row_index},"YYYY-MM")',
            23: f'=IF(L{row_index}="","",IF(L{row_index}>pLateDays,1,0))',
            24: f"=IF(Q{row_index}>=pHighCutoff,1,0)",
            25: f'=IF(M{row_index}="Returned",1,0)',
        }
        for col, formula in formulas.items():
            cell = ws.cell(row_index, col, formula)
            cell.font = BODY_FONT
            cell.border = THIN
        ws.cell(row_index, 16).number_format = '"$"#,##0.00'
        ws.cell(row_index, 17).number_format = '"$"#,##0.00'
        ws.cell(row_index, 18).number_format = '"$"#,##0.00'
        ws.cell(row_index, 19).number_format = "0.00%"
    last_row = len(export) + 1
    table = Table(displayName="CleanData", ref=f"A1:Y{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:Y{last_row}"
    for col in range(1, 26):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["F"].width = 18
    ws.sheet_properties.tabColor = TEAL
    ws.conditional_formatting.add(
        f"S2:S{last_row}",
        CellIsRule(operator="lessThan", formula=["pLowMargin"], fill=_fill(RED_SOFT)),
    )
    ws.conditional_formatting.add(
        f"X2:X{last_row}",
        CellIsRule(operator="equal", formula=["1"], fill=_fill(AMBER)),
    )


def _chart_data(wb, summary):
    ws = wb.create_sheet("Chart_Data")
    ws["A1"] = "Chart source — every number is a formula against CleanData and the dashboard parameters."
    ws["A1"].font = Font(name="Calibri", size=12, bold=True, color=NAVY)

    ws["A3"] = "Region"
    ws["B3"] = "NetRevenue"
    ws["C3"] = "Orders"
    ws["D3"] = "AvgMargin"
    ws["E3"] = "TotalCost"
    for offset, region in enumerate(REGIONS):
        row = 4 + offset
        ws.cell(row, 1, region)
        ws.cell(row, 2, sumifs("NetRevenue", honor_region=False, extras=(("Region", f"A{row}"),)))
        ws.cell(row, 3, countifs(honor_region=False, extras=(("Region", f"A{row}"),)))
        ws.cell(row, 4, averageifs("Margin", honor_region=False, extras=(("Region", f"A{row}"),)))
        ws.cell(row, 5, sumifs("TotalCost", honor_region=False, extras=(("Region", f"A{row}"),)))
        ws.cell(row, 2).number_format = '"$"#,##0'
        ws.cell(row, 4).number_format = "0.0%"
        ws.cell(row, 5).number_format = '"$"#,##0'

    ws["A9"] = "Channel"
    ws["B9"] = "Orders"
    ws["C9"] = "NetRevenue"
    for offset, channel in enumerate(CHANNELS):
        row = 10 + offset
        ws.cell(row, 1, channel)
        ws.cell(row, 2, countifs(honor_channel=False, extras=(("Channel", f"A{row}"),)))
        ws.cell(row, 3, sumifs("NetRevenue", honor_channel=False, extras=(("Channel", f"A{row}"),)))
        ws.cell(row, 3).number_format = '"$"#,##0'

    ws["A14"] = "YearMonth"
    ws["B14"] = "NetRevenue"
    ws["C14"] = "Orders"
    for offset, period in enumerate(summary["monthly"]["YearMonth"].tolist()):
        row = 15 + offset
        ws.cell(row, 1, period)
        ws.cell(row, 2, sumifs("NetRevenue", extras=(("YearMonth", f"A{row}"),)))
        ws.cell(row, 3, countifs(extras=(("YearMonth", f"A{row}"),)))
        ws.cell(row, 2).number_format = '"$"#,##0'
    month_end = 14 + len(summary["monthly"])

    ws["A37"] = "Category"
    for col, region in enumerate(REGIONS, start=2):
        ws.cell(37, col, region)
    for offset, category in enumerate(CATEGORIES):
        row = 38 + offset
        ws.cell(row, 1, category)
        for col, region in enumerate(REGIONS, start=2):
            letter = get_column_letter(col)
            formula = sumifs(
                "NetRevenue",
                honor_region=False,
                honor_category=False,
                extras=(("Category", f"A{row}"), ("Region", f"{letter}$37")),
            )
            cell = ws.cell(row, col, formula)
            cell.number_format = '"$"#,##0'

    ws["A43"] = "Quarter"
    ws["B43"] = "NetRevenue"
    for offset, quarter in enumerate(summary["quarterly"]["Quarter"].tolist()):
        row = 44 + offset
        ws.cell(row, 1, quarter)
        ws.cell(row, 2, sumifs("NetRevenue", extras=(("Quarter", f"A{row}"),)))
        ws.cell(row, 2).number_format = '"$"#,##0'
    quarter_end = 43 + len(summary["quarterly"])

    ws["A52"] = "Status"
    ws["B52"] = "Orders"
    for offset, status in enumerate(STATUSES):
        row = 53 + offset
        ws.cell(row, 1, status)
        ws.cell(row, 2, countifs_status(f"A{row}"))

    ws["A57"] = "Category"
    ws["B57"] = "ReturnRate"
    ws["C57"] = "Returned"
    ws["D57"] = "ReturnBase"
    for offset, category in enumerate(CATEGORIES):
        row = 58 + offset
        ws.cell(row, 1, category)
        ws.cell(row, 3, countifs_status('"Returned"', honor_category=False, extras=(("Category", f"A{row}"),)))
        ws.cell(row, 4, countifs_status('"<>Cancelled"', honor_category=False, extras=(("Category", f"A{row}"),)))
        ws.cell(row, 2, f"=IF(D{row}=0,NA(),C{row}/D{row})")
        ws.cell(row, 2).number_format = "0.0%"

    ws["A63"] = "YearMonth"
    for col, region in enumerate(REGIONS, start=2):
        ws.cell(63, col, region)
    for offset, period in enumerate(summary["monthly"]["YearMonth"].tolist()):
        row = 64 + offset
        ws.cell(row, 1, period)
        for col in range(2, 6):
            letter = get_column_letter(col)
            ws.cell(
                row,
                col,
                countifs(
                    honor_region=False,
                    extras=(("YearMonth", f"$A{row}"), ("Region", f"{letter}$63")),
                ),
            )
    region_month_end = 63 + len(summary["monthly"])

    ws["A86"] = "Region"
    for col, channel in enumerate(CHANNELS, start=2):
        ws.cell(86, col, channel)
    for offset, region in enumerate(REGIONS):
        row = 87 + offset
        ws.cell(row, 1, region)
        for col in range(2, 5):
            letter = get_column_letter(col)
            net = sumifs(
                "NetRevenue",
                honor_region=False,
                honor_channel=False,
                extras=(("Region", f"$A{row}"), ("Channel", f"{letter}$86")),
            )
            orders = countifs(
                honor_region=False,
                honor_channel=False,
                extras=(("Region", f"$A{row}"), ("Channel", f"{letter}$86")),
            )
            # sumifs/countifs already include the leading equals sign. Build the ratio from their bodies.
            ws.cell(row, col, f"=IF(({orders[1:]})=0,NA(),({net[1:]})/({orders[1:]}))")
            ws.cell(row, col).number_format = '"$"#,##0'

    for row in (3, 9, 14, 37, 43, 52, 57, 63, 86):
        for col in range(1, 6):
            if ws.cell(row, col).value:
                _paint(ws.cell(row, col), _fill(NAVY), HEADER_FONT, CENTER)
    ws.column_dimensions["A"].width = 16
    for col in range(2, 6):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.sheet_properties.tabColor = GOLD
    ws.sheet_view.showGridLines = False
    wb._chart_extents = {
        "month_end": month_end,
        "quarter_end": quarter_end,
        "region_month_end": region_month_end,
    }


def _analysis(wb, summary):
    ws = wb.create_sheet("Analysis")
    _banner(ws, 8, "Analysis", "Live SUMIFS summaries follow the dashboard filters. The top-10 list is a snapshot from the last pipeline run at All / All / All.")
    headers = ["Region", "Net revenue", "Orders", "Average order value", "Weighted margin", "Unweighted margin"]
    for col, header in enumerate(headers, start=2):
        _paint(ws.cell(4, col), _fill(NAVY), HEADER_FONT, CENTER, border=THIN)
    for offset, region in enumerate(REGIONS):
        row = 5 + offset
        source = 4 + offset
        ws.cell(row, 2, f"=Chart_Data!A{source}")
        ws.cell(row, 3, f"=Chart_Data!B{source}")
        ws.cell(row, 4, f"=Chart_Data!C{source}")
        ws.cell(row, 5, f"=IF(D{row}=0,NA(),C{row}/D{row})")
        ws.cell(row, 6, f"=IF(C{row}=0,NA(),(C{row}-Chart_Data!E{source})/C{row})")
        ws.cell(row, 7, f"=Chart_Data!D{source}")
        ws.cell(row, 3).number_format = '"$"#,##0'
        ws.cell(row, 5).number_format = '"$"#,##0'
        ws.cell(row, 6).number_format = "0.0%"
        ws.cell(row, 7).number_format = "0.0%"
        for col in range(2, 8):
            ws.cell(row, col).font = BODY_FONT
            ws.cell(row, col).border = THIN
            ws.cell(row, col).fill = _fill(WHITE)
    ws.conditional_formatting.add("G5:G8", CellIsRule(operator="lessThan", formula=["pLowMargin"], fill=_fill(RED_SOFT)))

    ws["B11"] = "Top 10 orders by net revenue (pipeline snapshot, default filters)"
    ws["B11"].font = Font(name="Calibri", size=13, bold=True, color=NAVY)
    top_headers = ["OrderID", "Order date", "Region", "Category", "Product", "Status", "Net revenue", "Margin"]
    for col, header in enumerate(top_headers, start=2):
        _paint(ws.cell(12, col), _fill(TEAL), HEADER_FONT, CENTER, border=THIN)
    top = summary["top_orders"]
    for offset, record in enumerate(top.itertuples(index=False)):
        row = 13 + offset
        values = [
            record.OrderID,
            record.OrderDate.to_pydatetime(),
            record.Region,
            record.Category,
            record.Product,
            record.Status,
            float(record.NetRevenue),
            None if pd.isna(record.Margin) else float(record.Margin),
        ]
        for col, value in enumerate(values, start=2):
            cell = ws.cell(row, col, value)
            cell.font = BODY_FONT
            cell.border = THIN
        ws.cell(row, 3).number_format = "YYYY-MM-DD"
        ws.cell(row, 8).number_format = '"$"#,##0.00'
        ws.cell(row, 9).number_format = "0.0%"
    note = (
        "Weighted margin is (net revenue - cost) / net revenue. Unweighted margin is the average of order-level margins. "
        "A region can look healthier on one definition than the other when a few large wholesale tickets dominate."
    )
    ws.merge_cells("B24:I26")
    ws["B24"] = note
    ws["B24"].alignment = WRAP
    ws["B24"].font = BODY_FONT
    for col, width in enumerate([3, 16, 22, 16, 18, 20, 16, 18, 14], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY


def _dashboard(wb, filters, params):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 4
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.sheet_view.zoomScale = 90
    last_col = 16
    _banner(ws, last_col, "Harbor & Co.  ·  Sales performance", "Excel pipeline dashboard. Change the controls, then let Excel calculate. Comparison charts ignore the filter that would collapse them.")
    labels = [
        (2, "Region"),
        (4, "Channel"),
        (6, "Category"),
        (8, "Include cancelled"),
        (10, "Start"),
        (12, "End"),
    ]
    for col, label in labels:
        _paint(ws.cell(4, col), _fill(CREAM), LABEL_FONT, CENTER)
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
    defaults = {
        3: filters["region"],
        5: filters["channel"],
        7: filters["category"],
        9: "Yes" if params["include_cancelled"] else "No",
        11: datetime.strptime(params["analysis_start"], "%Y-%m-%d"),
        13: datetime.strptime(params["analysis_end"], "%Y-%m-%d"),
    }
    for col, value in defaults.items():
        cell = ws.cell(5, col, value)
        _paint(cell, _fill(WHITE), Font(name="Calibri", size=12, bold=True, color=INK), CENTER, border=THIN)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        ws.cell(5, col + 1).border = THIN
        ws.cell(5, col + 1).fill = _fill(WHITE)
    ws["K5"].number_format = "YYYY-MM-DD"
    ws["M5"].number_format = "YYYY-MM-DD"
    ws.row_dimensions[5].height = 24
    note = "Shipment lateness and the low-margin floor live on Parameters. The high-value cutoff is calculated for this saved window; rerun the pipeline after changing dates if you need a new percentile."
    ws.merge_cells("A6:P6")
    ws["A6"] = note
    ws["A6"].font = Font(name="Calibri", size=9, italic=True, color=SLATE)
    ws["A6"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[6].height = 28

    wb.defined_names.add(DefinedName(name="pRegion", attr_text="Dashboard!$C$5"))
    wb.defined_names.add(DefinedName(name="pChannel", attr_text="Dashboard!$E$5"))
    wb.defined_names.add(DefinedName(name="pCategory", attr_text="Dashboard!$G$5"))
    wb.defined_names.add(DefinedName(name="pIncludeCancelled", attr_text="Dashboard!$I$5"))
    wb.defined_names.add(DefinedName(name="pStart", attr_text="Dashboard!$K$5"))
    wb.defined_names.add(DefinedName(name="pEnd", attr_text="Dashboard!$M$5"))

    region_dv = DataValidation(type="list", formula1="=lstRegion", allow_blank=False)
    channel_dv = DataValidation(type="list", formula1="=lstChannel", allow_blank=False)
    category_dv = DataValidation(type="list", formula1="=lstCategory", allow_blank=False)
    yes_dv = DataValidation(type="list", formula1="=lstYesNo", allow_blank=False)
    for dv in (region_dv, channel_dv, category_dv, yes_dv):
        dv.error = "Choose a listed value"
        dv.errorTitle = "Filter"
        dv.prompt = "Select a filter"
        dv.promptTitle = "Dashboard filter"
        ws.add_data_validation(dv)
    region_dv.add(ws["C5"])
    channel_dv.add(ws["E5"])
    category_dv.add(ws["G5"])
    yes_dv.add(ws["I5"])

    cards = [
        (1, "Net revenue", sumifs("NetRevenue"), '"$"#,##0', PALE),
        (3, "Orders", countifs(), "#,##0", CREAM),
        (5, "Average order value", f"=IF(C9=0,NA(),A9/C9)", '"$"#,##0', PALE),
        (7, "Gross margin", f"=IF(A9=0,NA(),(A9-{sumifs('TotalCost')[1:]})/A9)", "0.0%", AMBER),
        (9, "Return rate", f"=IF({countifs_status('\"<>Cancelled\"')[1:]}=0,NA(),{countifs_status('\"Returned\"')[1:]}/{countifs_status('\"<>Cancelled\"')[1:]})", "0.0%", PALE),
        (11, "On-time fulfillment", f"=IF({countifs_status('\"<>Cancelled\"', extras=(('ShipDays', '\">=\"&0'),))[1:]}=0,NA(),{countifs_status('\"<>Cancelled\"', extras=(('LateFlag', '0'),))[1:]}/{countifs_status('\"<>Cancelled\"', extras=(('ShipDays', '\">=\"&0'),))[1:]})", "0.0%", CREAM),
        (13, "High-value orders", countifs(extras=(("HighValueFlag", "1"),)), "#,##0", AMBER),
        (15, "Low-margin orders", countifs(extras=(("Margin", '"<"&pLowMargin'),)), "#,##0", RED_SOFT),
    ]
    for col, label, formula, fmt, color in cards:
        ws.merge_cells(start_row=8, start_column=col, end_row=8, end_column=col + 1)
        ws.merge_cells(start_row=9, start_column=col, end_row=9, end_column=col + 1)
        label_cell = ws.cell(8, col, label)
        value_cell = ws.cell(9, col, formula)
        _paint(label_cell, _fill(color), LABEL_FONT, CENTER)
        _paint(value_cell, _fill(color), KPI_FONT, CENTER, fmt)
        ws.cell(8, col + 1).fill = _fill(color)
        ws.cell(9, col + 1).fill = _fill(color)
    ws.row_dimensions[8].height = 18
    ws.row_dimensions[9].height = 36

    extents = wb._chart_extents
    specs = [
        ("Net revenue by region", "col", "A4:A7", "B3:B7", "A11"),
        ("Orders by channel", "pie", "A10:A12", "B9:B12", "I11"),
        ("Monthly net revenue", "line", f"A15:A{extents['month_end']}", f"B14:B{extents['month_end']}", "A27"),
        ("Net revenue by category and region", "stack", "A38:A41", "B37:E41", "I27"),
        ("Unweighted average margin by region", "col", "A4:A7", "D3:D7", "A43"),
        ("Quarterly net revenue", "area", f"A44:A{extents['quarter_end']}", f"B43:B{extents['quarter_end']}", "I43"),
        ("Order status mix", "donut", "A53:A55", "B52:B55", "A59"),
        ("Return rate by category", "col", "A58:A61", "B57:B61", "I59"),
        ("Monthly orders by region", "line4", f"A64:A{extents['region_month_end']}", f"B63:E{extents['region_month_end']}", "A75"),
        ("Average order value by region and channel", "cluster", "A87:A90", "B86:D90", "I75"),
    ]
    for title, kind, cats, data, anchor in specs:
        _add_dashboard_chart(ws, title, kind, cats, data, anchor)

    for col in range(1, 17):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.row_dimensions[4].height = 18
    ws.page_setup.horizontalCentered = True


def _add_dashboard_chart(ws, title, kind, cats, data, anchor):
    cat_ref = _ref("Chart_Data", cats)
    data_ref = _ref("Chart_Data", data)
    if kind == "pie":
        chart = PieChart()
        chart.title = title
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showVal = False
        chart.dataLabels.showCatName = False
        _apply_points(chart, SERIES)
    elif kind == "donut":
        chart = DoughnutChart()
        chart.title = title
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        _apply_points(chart, [TEAL, GOLD, CORAL])
    elif kind == "area":
        chart = AreaChart()
        chart.title = title
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        chart.y_axis.numFmt = '"$"#,##0'
        _color_series(chart, [TEAL])
    elif kind == "line":
        chart = LineChart()
        chart.title = title
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        chart.y_axis.numFmt = '"$"#,##0'
        _color_series(chart, [NAVY])
    elif kind == "line4":
        chart = LineChart()
        chart.title = title
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        chart.y_axis.numFmt = "#,##0"
        _legend(chart)
        _color_series(chart)
    elif kind == "stack":
        chart = _bar(title, ws, data_ref, cat_ref, anchor, stacked=True)
        return chart
    elif kind == "cluster":
        chart = _bar(title, ws, data_ref, cat_ref, anchor, stacked=False)
        chart.y_axis.numFmt = '"$"#,##0'
        return chart
    else:
        chart = _bar(title, ws, data_ref, cat_ref, anchor, stacked=False, colors=[NAVY])
        if "margin" in title.lower() or "return" in title.lower():
            chart.y_axis.numFmt = "0%"
        return chart
    _legend(chart)
    _size(chart)
    ws.add_chart(chart, anchor)
    return chart


def _ref(sheet, a1_range):
    start, end = a1_range.split(":")
    start_col = "".join(ch for ch in start if ch.isalpha())
    start_row = int("".join(ch for ch in start if ch.isdigit()))
    end_col = "".join(ch for ch in end if ch.isalpha())
    end_row = int("".join(ch for ch in end if ch.isdigit()))
    return Reference(
        None,
        min_col=_col_index(start_col),
        min_row=start_row,
        max_col=_col_index(end_col),
        max_row=end_row,
    ) if False else _reference_with_sheet(sheet, start_col, start_row, end_col, end_row)


def _col_index(letters: str) -> int:
    value = 0
    for char in letters:
        value = value * 26 + (ord(char.upper()) - 64)
    return value


def _reference_with_sheet(sheet, start_col, start_row, end_col, end_row):
    # Reference needs a worksheet for some openpyxl versions; a range string is enough
    # when the chart lives on another sheet if we set the title via the formula.
    ref = Reference(
        range_string=f"{sheet}!${start_col}${start_row}:${end_col}${end_row}"
    )
    return ref


def _apply_points(chart, colors):
    if not chart.series:
        return
    points = []
    for index, color in enumerate(colors):
        point = DataPoint(idx=index)
        point.graphicalProperties = GraphicalProperties(solidFill=color)
        points.append(point)
    chart.series[0].data_points = points


def _insights(wb, summary, params, cleaning_log):
    ws = wb.create_sheet("Insights")
    _banner(ws, 6, "Insights", "Written from the latest pipeline run at the default filters.")
    narrative = summary.get("narrative", "")
    ws.merge_cells("B4:F28")
    ws["B4"] = narrative
    ws["B4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B4"].font = Font(name="Calibri", size=12, color=INK)
    ws.column_dimensions["B"].width = 28
    for col in range(3, 7):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.row_dimensions[4].height = 320
    ws.sheet_view.showGridLines = False


def _qa(wb, summary, cleaning_log, clean, raw, rejected):
    ws = wb.create_sheet("QA_Tests")
    _banner(ws, 6, "QA tests", "Status formulas compare Excel calculations with the values locked by the pipeline. Open in Excel and calculate.")
    headers = ["Test", "Expected", "Actual (formula)", "Status", "Notes"]
    for col, header in enumerate(headers, start=2):
        _paint(ws.cell(4, col), _fill(NAVY), HEADER_FONT, CENTER, border=THIN)
    cards = summary["kpis"]
    north = float(summary["regions"].loc[summary["regions"]["Name"] == "North", "NetRevenue"].iloc[0])
    tests = [
        ("Raw rows", len(raw), "=COUNTA(RawData[OrderID])", "Every source row, including duplicates and edge cases."),
        ("Clean rows", len(clean), "=COUNTA(CleanData[OrderID])", "One row per surviving OrderID."),
        ("Rejected rows", len(rejected), "=COUNTA(RejectedData[OrderID])", "Duplicates, invalid rows, and OrderID conflicts."),
        ("Row balance", len(raw), "=COUNTA(CleanData[OrderID])+COUNTA(RejectedData[OrderID])", "Raw count must equal clean plus rejected."),
        ("Net revenue", round(cards["net_revenue"], 2), "=ROUND(" + sumifs("NetRevenue")[1:] + ",2)", "Default filters, cancelled excluded."),
        ("Orders", cards["orders"], countifs(), "Default filters, cancelled excluded."),
        ("North net revenue", round(north, 2), "=ROUND(Chart_Data!B4,2)", "Region chart ignores the region dropdown and honors the other filters."),
        ("High-value cutoff", round(summary["cutoff"], 2), "=ROUND(pHighCutoff,2)", "Percentile cutoff stored at pipeline run time."),
        ("High-value orders", cards["high_value_orders"], countifs(extras=(("HighValueFlag", "1"),)), "Uses the stored cutoff."),
        ("Dashboard region default", "All", "=pRegion", "Named range points at the dashboard control."),
        ("Include cancelled default", "No" if not True else "Yes", "=pIncludeCancelled", "Sales KPIs exclude cancelled orders until this switch is Yes."),
    ]
    # Fix include cancelled expected from the real parameter via the formula default already written.
    tests[10] = (
        "Include cancelled default",
        "Yes" if summary.get("include_cancelled") else "No",
        "=pIncludeCancelled",
        "Sales KPIs exclude cancelled orders until this switch is Yes.",
    )
    for offset, (name, expected, actual, note) in enumerate(tests):
        row = 5 + offset
        ws.cell(row, 2, name)
        ws.cell(row, 3, expected)
        ws.cell(row, 4, actual if str(actual).startswith("=") else actual)
        ws.cell(row, 5, f'=IF(ABS(C{row}-D{row})<0.05,"PASS","FAIL")' if isinstance(expected, (int, float)) else f'=IF(C{row}=D{row},"PASS","FAIL")')
        ws.cell(row, 6, note)
        for col in range(2, 7):
            ws.cell(row, col).font = BODY_FONT
            ws.cell(row, col).border = THIN
            ws.cell(row, col).alignment = WRAP
        if isinstance(expected, float):
            ws.cell(row, 3).number_format = '#,##0.00'
        ws.row_dimensions[row].height = 32
    ws.conditional_formatting.add("E5:E15", FormulaRule(formula=['E5="PASS"'], fill=_fill("D9EFE8")))
    ws.conditional_formatting.add("E5:E15", FormulaRule(formula=['E5="FAIL"'], fill=_fill(RED_SOFT)))
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 78
    ws.sheet_view.showGridLines = False
    # Rejected rows table so COUNTA(RejectedData[OrderID]) resolves.
    rejected_sheet = wb.create_sheet("Rejected_Rows")
    export = rejected.copy()
    if export.empty:
        export = pd.DataFrame([{"OrderID": None, "Reason": None}])
    _write_table(rejected_sheet, export.fillna(""), "RejectedData")
    rejected_sheet.sheet_properties.tabColor = CORAL


def _dictionary(wb):
    ws = wb.create_sheet("Data_Dictionary")
    _banner(ws, 4, "Data dictionary", "Field definitions for the cleaned orders table.")
    headers = ["Field", "Meaning", "Notes"]
    for col, header in enumerate(headers, start=2):
        _paint(ws.cell(4, col), _fill(NAVY), HEADER_FONT, CENTER, border=THIN)
    rows = [
        ("OrderID", "Order identifier", "Conflicts keep the higher net-revenue row."),
        ("OrderDate", "Order date", "Parsed from mixed raw formats, then stored as a real date."),
        ("Region", "Sales region", "North, South, East, West after alias cleanup."),
        ("Channel", "Route to customer", "Store, Online, or Wholesale."),
        ("Category", "Merchandise family", "Furniture, Lighting, Textiles, Kitchen."),
        ("Product", "SKU name", "Free text from the catalog."),
        ("Segment", "Customer segment", "Blank source values become Unknown."),
        ("Quantity", "Units", "Negative quantities are flipped and flagged. Zero quantities are rejected."),
        ("UnitPrice", "Selling price per unit", "Currency text and spelled numbers are parsed in the transform."),
        ("Discount", "Discount fraction", "20% and 0.20 both become 0.20. Values between 1 and 100 are treated as percents."),
        ("UnitCost", "Cost per unit", "Must be zero or positive."),
        ("ShipDays", "Days from order to ship", "Blank, N/A, negative, or above 120 days are cleared and excluded from on-time rate."),
        ("Status", "Order outcome", "Completed, Returned, Cancelled."),
        ("GrossRevenue", "Quantity x unit price", "Excel formula on Cleaned_Data."),
        ("NetRevenue", "Gross x (1 - discount)", "Excel formula. This is the sales KPI numerator."),
        ("TotalCost", "Quantity x unit cost", "Excel formula."),
        ("Margin", "(Net - cost) / net", "Blank when net revenue is zero. This is an unweighted order margin."),
        ("LateFlag", "1 when ShipDays exceeds the parameter", "Blank when ship days are missing."),
        ("HighValueFlag", "1 when net revenue >= cutoff", "Cutoff is the pipeline percentile, not a live percentile of the filtered view."),
        ("ReturnFlag", "1 when status is Returned", "Used for audit, not for the return-rate denominator."),
    ]
    for offset, values in enumerate(rows):
        for col, value in enumerate(values, start=2):
            cell = ws.cell(5 + offset, col, value)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = THIN
        ws.row_dimensions[5 + offset].height = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 48
    ws.column_dimensions["D"].width = 78
    ws.sheet_view.showGridLines = False


def _cleaning_log(wb, cleaning_log):
    ws = wb.create_sheet("Cleaning_Log")
    _banner(ws, 4, "Cleaning log", "Counts from the transform that produced this workbook.")
    ws["B4"] = "Metric"
    ws["C4"] = "Value"
    _paint(ws["B4"], _fill(NAVY), HEADER_FONT, CENTER)
    _paint(ws["C4"], _fill(NAVY), HEADER_FONT, CENTER)
    rows = [
        ("Raw rows", cleaning_log["raw_rows"]),
        ("Clean rows", cleaning_log["clean_rows"]),
        ("Rejected rows", cleaning_log["rejected_rows"]),
        ("Exact duplicates", cleaning_log["exact_duplicates"]),
        ("Row balance holds", "Yes" if cleaning_log["balance_ok"] else "No"),
    ]
    for reason, count in sorted(cleaning_log["reason_counts"].items()):
        rows.append((f"Rejected: {reason}", count))
    for name, count in sorted(cleaning_log["adjustments"].items()):
        rows.append((f"Adjustment: {name}", count))
    for offset, (label, value) in enumerate(rows):
        ws.cell(5 + offset, 2, label).font = BODY_FONT
        ws.cell(5 + offset, 3, value).font = BODY_FONT
        ws.cell(5 + offset, 2).border = THIN
        ws.cell(5 + offset, 3).border = THIN
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 18
    ws.sheet_view.showGridLines = False


def _cover(wb, params, cleaning_log, summary):
    ws = wb.create_sheet("Cover")
    _banner(ws, 6, "Harbor & Co. advanced data analysis", "Excel pipeline  ·  extract, clean, analyze, dashboard, test")
    cards = summary["kpis"]
    blocks = [
        (4, "Objective", "Turn a dirty order extract into a decision-ready sales view: what sold, where, through which channel, and which orders are late, returned, or thin on margin."),
        (8, "Pipeline", "1. Raw_Data holds the dirty extract.\n2. pipeline/clean.py standardizes aliases, dates, numbers, duplicates, and invalid rows.\n3. Cleaned_Data stores the surviving orders. Revenue, margin, and flags are Excel formulas.\n4. Chart_Data and Dashboard aggregate with SUMIFS / COUNTIFS / AVERAGEIFS so filters recalculate in Excel.\n5. QA_Tests checks those formulas against the values locked by the Python run."),
        (13, "Default window", f"{params['analysis_start']} to {params['analysis_end']}. Cancelled orders are excluded from sales KPIs. Seed {params['random_seed']}. Raw rows {cleaning_log['raw_rows']:,}. Clean rows {cleaning_log['clean_rows']:,}."),
        (16, "Headline result", f"Net revenue ${cards['net_revenue']:,.0f}  ·  Orders {cards['orders']:,}  ·  AOV ${cards['aov']:,.0f}  ·  Gross margin {cards['gross_margin']*100:.1f}%  ·  Return rate {cards['return_rate']*100:.1f}%  ·  On time {cards['on_time_rate']*100:.1f}%."),
        (19, "How to use the dashboard", "Open Dashboard. Pick Region, Channel, Category, the cancelled switch, and the date window. Excel recalculates the KPI cards and all ten charts. Comparison charts keep their breakdown: the region chart does not collapse when you pick one region."),
    ]
    for row, title, body in blocks:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        _paint(ws.cell(row, 2, title), _fill(TEAL), HEADER_FONT, Alignment(vertical="center", indent=1))
        for col in range(2, 7):
            ws.cell(row, col).fill = _fill(TEAL)
        ws.merge_cells(start_row=row + 1, start_column=2, end_row=row + 2, end_column=6)
        cell = ws.cell(row + 1, 2, body)
        cell.alignment = WRAP
        cell.font = BODY_FONT
        ws.row_dimensions[row + 1].height = 22
        ws.row_dimensions[row + 2].height = 36
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 24
    for col in range(3, 7):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
